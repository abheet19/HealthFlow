"""Report-content regression: run from backend with python -m unittest discover -s tests."""
import unittest
from pathlib import Path
from docx import Document
from app.services.report_service import ReportService

class ReportTests(unittest.TestCase):
    def test_vision_acuity_is_not_color_blindness(self):
        record={"name":"Synthetic Report Test", "rev":"6/6", "lev":"6/9", "rcb":"No", "lcb":"Yes", "photo":None}
        report,_=ReportService.generate_word_report(None,record,str(Path(__file__).resolve().parents[1]/"template.docx"))
        doc=Document(report)
        text=" ".join(p.text for p in doc.paragraphs)+" ".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        vision=text.split("VISION",1)[1].split("COLOR BLINDNESS",1)[0]
        self.assertIn("6/6",vision)
        self.assertIn("6/9",vision)
        self.assertNotIn("{{",text)
        color=text.split("COLOR BLINDNESS",1)[1].split("SQUINT",1)[0]
        self.assertIn("No",color)
        self.assertIn("Yes",color)

    def test_corrupt_photo_does_not_prevent_report(self):
        report,_=ReportService.generate_word_report(None,{"name":"Synthetic Test", "photo":"not-base64!"},str(Path(__file__).resolve().parents[1]/"template.docx"))
        self.assertGreater(len(report.getvalue()),1000)

if __name__ == "__main__": unittest.main()
